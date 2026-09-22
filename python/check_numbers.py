# -*- coding: utf-8 -*-
"""科研证据链校验器（通用实现，随 dsh-research-check 插件分发）。

设计目标：把“论文里的数字、图注里的数字、结果文件里的数字”三者对齐，
并检查投稿硬约束与文件卫生。所有检查器都只读，输出单个 JSON 报告。

用法（由插件调用，也可单独运行）：
  python check_numbers.py --paper a.pdf b.tex --ledger ledger.json [--tol 0.01]
  python audit_paper.py   --paper a.pdf [--max-pages 30] [--extras x.zip]
  python check_hygiene.py --files a.xlsx b.docx [--json out.json]

数值抽取策略：
  · 文本源（.tex/.md/.txt）：去注释与数学环境标记后，抓取 R(v) 与“数值+单位”
  · PDF：逐页取文字层，同样抓取数值并记录页码
  · 表格源（.csv/.xlsx）：抓取全部数值单元格（作为“数据侧”候选）
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys
from pathlib import Path
from encoding_guard import force_utf8_output


# --------------------------------------------------------------------------
# 数值抽取
# --------------------------------------------------------------------------
# 只保留“业务量纲”。刻意排除 s/m/h 这类短单位：代码附录与行号会把
# “数字+s”误判为秒，产生大量噪声（实测 5 条警告全部来自误判）。
NUM_UNITS = (
    "kWh", "MWh", "GWh", "kW", "MW", "GW", "元", "万元", "亿元",
    "小时", "分钟", "%", "美元", "度", "件", "天", "台",
)
UNIT_ALT = "|".join(re.escape(u) for u in sorted(NUM_UNITS, key=len, reverse=True))
NUM = r"[-+]?\d[\d,]*(?:\.\d+)?(?:[eE][-+]?\d+)?"
PAT_UNIT = re.compile(rf"({NUM})\s*({UNIT_ALT})")
PAT_PLAIN = re.compile(rf"(?<![\d.])({NUM})(?![\d.])")


def strip_latex(text: str) -> str:
    """Remove comments and math delimiters while keeping the numbers intact."""
    text = re.sub(r"(?<!\\)%.*", "", text)
    text = re.sub(r"\\(?:label|ref|eqref|cite|upcite|includegraphics)(?:\[[^\]]*\])?\{[^}]*\}", " ", text)
    text = text.replace("\\,", " ").replace("\\ ", " ").replace("~", " ")
    for token in ("$", "\\(", "\\)", "\\[", "\\]"):
        text = text.replace(token, " ")
    text = re.sub(r"\\[a-zA-Z]+\*?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return text


def extract_from_text(text: str, source: str, page: int | None = None) -> list[dict]:
    """Extract candidate numbers with their nearest unit and local context."""
    hits: list[dict] = []
    cleaned = strip_latex(text)
    for match in PAT_UNIT.finditer(cleaned):
        hits.append(make_hit(match, cleaned, source, page))
    for match in PAT_PLAIN.finditer(cleaned):
        hit = make_hit(match, cleaned, source, page)
        hit["has_unit"] = False
        hits.append(hit)
    return hits


def make_hit(match: re.Match, cleaned: str, source: str, page: int | None) -> dict:
    raw = match.group(1)
    unit = match.group(2) if match.lastindex and match.lastindex >= 2 else None
    start, end = max(0, match.start() - 46), min(len(cleaned), match.end() + 46)
    return {
        "value": to_float(raw),
        "raw": raw,
        "unit": unit,
        "has_unit": unit is not None,
        "source": source,
        "page": page,
        "context": re.sub(r"\s+", " ", cleaned[start:end]).strip(),
    }


def to_float(raw: str) -> float | None:
    try:
        return float(raw.replace(",", ""))
    except ValueError:
        return None


def read_pdf(path: Path) -> list[dict]:
    try:
        import pymupdf  # type: ignore
    except ImportError:
        return []
    hits: list[dict] = []
    doc = pymupdf.open(path)
    for index, page in enumerate(doc, start=1):
        hits.extend(extract_from_text(page.get_text(), path.name, index))
    doc.close()
    return hits


def read_text_file(path: Path) -> list[dict]:
    return extract_from_text(path.read_text(encoding="utf-8", errors="ignore"), path.name)


def read_table_file(path: Path) -> list[dict]:
    """Numbers coming from data files (the 'evidence' side)."""
    hits: list[dict] = []
    if path.suffix.lower() == ".csv":
        for line_no, line in enumerate(path.read_text(encoding="utf-8", errors="ignore").splitlines(), 1):
            for cell in line.split(","):
                value = to_float(cell.strip())
                if value is not None:
                    hits.append({
                        "value": value, "raw": cell.strip(), "unit": None, "has_unit": False,
                        "source": path.name, "page": None, "context": f"第 {line_no} 行",
                    })
        return hits
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            import openpyxl  # type: ignore
        except ImportError:
            return hits
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for row_index, row in enumerate(ws.iter_rows(values_only=True), 1):
                for col_index, cell in enumerate(row, 1):
                    if isinstance(cell, (int, float)) and not isinstance(cell, bool):
                        hits.append({
                            "value": float(cell), "raw": repr(cell), "unit": None,
                            "has_unit": False, "source": f"{path.name}!{sheet}",
                            "page": None, "context": f"r{row_index}c{col_index}",
                        })
        wb.close()
        return hits
    return hits


def load_document(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix in {".tex", ".md", ".txt", ".rst"}:
        return read_text_file(path)
    if suffix in {".csv", ".xlsx", ".xlsm"}:
        return read_table_file(path)
    return []


# --------------------------------------------------------------------------
# 比对
# --------------------------------------------------------------------------
def close(a: float, b: float, rel: float, abs_tol: float) -> bool:
    scale = max(abs(a), abs(b), 1e-12)
    return abs(a - b) <= max(abs_tol, rel * scale)


def split_documents(paths: list[Path]) -> tuple[list[Path], list[Path]]:
    """Paper-side sources vs data-side sources."""
    papers, data = [], []
    for path in paths:
        if path.suffix.lower() in {".csv", ".xlsx", ".xlsm"}:
            data.append(path)
        else:
            papers.append(path)
    return papers, data


def verify_ledger(ledger: dict, paper_hits: list[dict], data_hits: list[dict],
                  rel: float, abs_tol: float) -> list[dict]:
    findings: list[dict] = []
    entries = ledger.get("entries", []) if isinstance(ledger, dict) else []
    for entry in entries:
        key = entry.get("key", "?")
        expected = entry.get("value")
        if not isinstance(expected, (int, float)):
            continue
        matched = [hit for hit in paper_hits if close(hit["value"], expected, rel, abs_tol)]
        if not matched:
            findings.append({
                "level": "error", "key": key, "expected": expected,
                "message": f"台账值 {expected}{entry.get('unit') or ''} 在论文中找不到；"
                           "可能是改写时漏改，或台账未随稿更新。",
                "where": entry.get("origin") or entry.get("note") or "",
            })
            continue
        # 允许同值多处出现；只在“带单位且单位一致”的命中里检查数值冲突
        unit = entry.get("unit")
        if unit:
            wrong_unit = [
                hit for hit in paper_hits
                if hit.get("unit") == unit and not close(hit["value"], expected, rel, abs_tol)
                and same_context(hit, entry)
            ]
            for hit in wrong_unit[:3]:
                findings.append({
                    "level": "warning", "key": key, "expected": expected, "found": hit["value"],
                    "message": f"同一位置疑似应为 {expected}{unit}，论文中为 {hit['value']}{unit}。",
                    "where": f"{hit['source']}"
                             + (f" 第 {hit['page']} 页" if hit.get("page") else "")
                             + f"｜{hit['context'][:60]}",
                })
    return findings


def same_context(hit: dict, entry: dict) -> bool:
    """Cheap anchor test: does the ledger entry name a keyword near the hit?"""
    anchors = [a for a in re.split(r"[｜|,，、;；\s]+", str(entry.get("anchor") or "")) if len(a) >= 2]
    if not anchors:
        return True
    return any(anchor in hit.get("context", "") for anchor in anchors)


def cross_source_conflicts(papers: list[Path], docs: dict[str, list[dict]],
                           rel: float, abs_tol: float) -> list[dict]:
    """Numbers carrying the same unit that appear as two different values in the
    same document are a strong self-contradiction signal."""
    findings: list[dict] = []
    for path in papers:
        hits = docs.get(str(path), [])
        buckets: dict[tuple[str, str], list[dict]] = {}
        for hit in hits:
            if not hit.get("unit"):
                continue
            if abs(hit["value"]) < 10:
                # 小数值（系数、指数、页码、行号）几乎都是噪声，不作为冲突证据
                continue
            anchor = re.sub(r"\d", "#", hit.get("context", ""))[:40]
            if anchor == "":
                continue
            buckets.setdefault((hit["unit"], anchor), []).append(hit)
        for (unit, anchor), group in buckets.items():
            values = {round(h["value"], 6) for h in group}
            if len(values) <= 1:
                continue
            ordered = sorted(values, key=abs, reverse=True)
            biggest = ordered[0]
            others = [v for v in ordered[1:] if not close(v, biggest, rel, abs_tol)]
            if not others:
                continue
            findings.append({
                "level": "warning",
                "message": f"同一句式的{unit}数值出现多个取值：{', '.join(f'{v:g}' for v in ordered[:4])}",
                "where": f"{path.name}｜…{anchor[:36]}…",
                "values": ordered[:6],
            })
    return findings


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------
def main(argv: list[str] | None = None) -> int:
    force_utf8_output()
    parser = argparse.ArgumentParser(description="论文数字一致性校验")
    parser.add_argument("--paper", nargs="+", required=True, help="论文与文本源（.pdf/.tex/.md/.txt）")
    parser.add_argument("--data", nargs="*", default=[], help="数据源（.csv/.xlsx），作为证据侧")
    parser.add_argument("--ledger", help="台账 JSON；给出后逐条核对")
    parser.add_argument("--tol", type=float, default=0.01, help="相对容差，默认 1%%")
    parser.add_argument("--abs-tol", type=float, default=0.005, help="绝对容差，默认 0.005")
    parser.add_argument("--max-findings", type=int, default=200)
    parser.add_argument("--emit-hits", action="store_true",
                        help="附带抽取到的数字清单（供台账工具 teach/verify 使用，输出较大）")
    parser.add_argument("--hits-limit", type=int, default=20000)
    args = parser.parse_args(argv)

    paper_paths = [Path(p) for p in args.paper]
    data_paths = [Path(p) for p in args.data]
    docs: dict[str, list[dict]] = {}
    for path in [*paper_paths, *data_paths]:
        if not path.exists():
            print(json.dumps({"ok": False, "error": "MISSING_INPUT", "path": str(path)},
                             ensure_ascii=False), flush=True)
            return 2
        docs[str(path)] = load_document(path)

    paper_hits = [h for p in paper_paths for h in docs.get(str(p), [])]
    data_hits = [h for p in data_paths for h in docs.get(str(p), [])]

    findings: list[dict] = []
    if args.ledger:
        ledger_path = Path(args.ledger)
        if not ledger_path.exists():
            print(json.dumps({"ok": False, "error": "MISSING_LEDGER", "path": str(ledger_path)},
                             ensure_ascii=False), flush=True)
            return 2
        ledger = json.loads(ledger_path.read_text(encoding="utf-8"))
        findings += verify_ledger(ledger, paper_hits, data_hits, args.tol, args.abs_tol)
    findings += cross_source_conflicts(paper_paths, docs, args.tol, args.abs_tol)

    order = {"error": 0, "warning": 1, "info": 2}
    findings.sort(key=lambda f: order.get(f.get("level", "info"), 3))
    counts = {"error": 0, "warning": 0, "info": 0}
    for finding in findings:
        counts[finding.get("level", "info")] = counts.get(finding.get("level", "info"), 0) + 1

    report = {
        "ok": True,
        "documents": len(paper_paths) + len(data_paths),
        "paper_numbers": len(paper_hits),
        "data_numbers": len(data_hits),
        "ledger_entries": len(json.loads(Path(args.ledger).read_text(encoding="utf-8")).get("entries", []))
        if args.ledger and Path(args.ledger).exists() else 0,
        "verdict": "fail" if counts["error"] else ("warn" if counts["warning"] else "pass"),
        "counts": counts,
        "findings": findings[: args.max_findings],
        "numbers": paper_hits[: args.hits_limit] if args.emit_hits else [],
    }
    print(json.dumps(report, ensure_ascii=False, indent=2), flush=True)
    return 1 if counts["error"] else 0


if __name__ == "__main__":
    sys.exit(main())
